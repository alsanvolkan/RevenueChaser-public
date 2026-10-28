"""``funds/index_holdings``: index membership from the ETF holdings files and Nasdaq's list API (daily).

**Not run against the real sources.** Each parser is built against the provider's published file format; the tests
use small files of that shape. The first real fetch is a homesrv step, and a format that has moved is a
``source_failed`` build issue for that file (the others go on), never a guess.

| Index | Files |
|---|---|
| ``sp500`` | SSGA SPY (xlsx), iShares IVV (csv), Vanguard VOO (json) |
| ``sp400`` | SSGA MDY, iShares IJH, Vanguard IVOO |
| ``sp600`` | SSGA SPSM, iShares IJR, Vanguard VIOO |
| ``ndx`` | Nasdaq ``/api/quote/list-type/nasdaq100`` (json) |

One request per file, every file each run (10 a run), each under its provider's day cap. Base URLs:
``RC_SSGA_BASE``, ``RC_ISHARES_BASE``, ``RC_VANGUARD_BASE``, ``RC_NASDAQ_API_BASE``.

**The membership rule (working choice).** On each day, a file counts when it parsed and holds at least its index's
floor (``FLOOR``: about 80% of the index); a short file is a ``fund_file_short`` build issue and does not count. A
ticker is a member of the index that day when more than half of the counting files hold it (one file of one: the
Nasdaq list). Cash, futures and other non-equity lines are dropped by each parser. Tickers are written as the SEC file
writes them (``BRK.B`` and ``BRK/B`` -> ``BRK-B``); ``spine/security`` matches them to the SEC file, by CUSIP first.

Raw: ``raw/funds/index_holdings/<as-of>/<provider>-<fund>.<ext>`` (+ ``.meta.json``). Parquet:
``parquet/funds/index_holdings/date=<as-of>/part-0.parquet`` (every holding line kept, with its file). Staging:
``stg_fund_lines`` (every day's lines) and ``stg_fund_files`` (which files counted each day).
"""
from __future__ import annotations

import csv
import io
import json
import os
import re
import time
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Callable

import duckdb

from ....core import calendar
from .. import base
from . import KIND, stage

NAME = "funds/index_holdings"
INDEXES = ("sp500", "sp400", "sp600", "ndx")
FLOOR = {"sp500": 400, "sp400": 320, "sp600": 480, "ndx": 80}
SCHEMA = [("as_of", "DATE"), ("index_slug", "TEXT"), ("provider", "TEXT"), ("fund", "TEXT"), ("ticker", "TEXT"),
          ("ticker_raw", "TEXT"), ("cusip", "TEXT"), ("name", "TEXT"), ("weight", "DOUBLE"),
          ("fetched_at", "DOUBLE"), ("run_id", "TEXT")]
FILE_SCHEMA = [("as_of", "DATE"), ("index_slug", "TEXT"), ("provider", "TEXT"), ("fund", "TEXT"),
               ("lines", "INTEGER"), ("counted", "BOOLEAN")]
META_SUFFIX = ".meta.json"
_TICKER = re.compile(r"^[A-Z][A-Z0-9]{0,5}(-[A-Z0-9]{1,3})?$")
# iShares' product pages, by fund: (product id, page slug). The CSV is the page's holdings download.
ISHARES_PRODUCTS = {"IVV": ("239726", "ishares-core-sp-500-etf"), "IJH": ("239763", "ishares-core-sp-midcap-etf"),
                    "IJR": ("239774", "ishares-core-sp-smallcap-etf")}


def norm_ticker(raw) -> str | None:
    """A holdings ticker as the SEC file writes it; None for a line that is no stock ticker (cash, futures, ``-``)."""
    t = str(raw or "").strip().upper().replace(".", "-").replace("/", "-").replace(" ", "-")
    return t if _TICKER.match(t) else None


def _cusip(raw) -> str | None:
    t = str(raw or "").strip().upper()
    return t if len(t) == 9 and t.isalnum() else None


def _weight(raw) -> float | None:
    try:
        return float(str(raw).replace("%", "").replace(",", "").strip())
    except (TypeError, ValueError):
        return None


# ── Parsers: bytes -> [{ticker_raw, cusip, name, weight}] ──────────────────────────────────────────────────────

def parse_ssga(data: bytes) -> list[dict]:
    """SSGA's daily holdings xlsx: preamble rows, then a header row ``Name, Ticker, Identifier, SEDOL, Weight,
    Sector, Shares Held, Local Currency``; lines until the first empty row. ``Identifier`` is the CUSIP."""
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
    head, out = None, []
    for row in ws.iter_rows(values_only=True):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if head is None:
            if "Ticker" in cells and "Name" in cells:
                head = {name: i for i, name in enumerate(cells) if name}
            continue
        if not any(cells):
            break
        get = lambda col: cells[head[col]] if col in head and head[col] < len(cells) else ""  # noqa: E731
        if get("Sector") in ("", "-") and not _cusip(get("Identifier")):
            continue                                         # cash and futures lines carry no sector and no CUSIP
        out.append({"ticker_raw": get("Ticker"), "cusip": _cusip(get("Identifier")), "name": get("Name"),
                    "weight": _weight(get("Weight"))})
    if head is None:
        raise base.SourceError("ssga: no header row with Name and Ticker")
    return out


def parse_ishares(data: bytes) -> list[dict]:
    """iShares' holdings CSV: preamble lines (fund name, ``Fund Holdings as of``), a header line starting
    ``Ticker,Name,Sector,Asset Class,...``, the lines, then a blank line and a disclaimer. Equity lines only."""
    text = data.decode("utf-8-sig", errors="replace")
    lines = text.splitlines()
    start = next((i for i, line in enumerate(lines) if line.startswith("Ticker,")), None)
    if start is None:
        raise base.SourceError("ishares: no header line starting with Ticker")
    body = []
    for line in lines[start:]:
        if not line.strip() or line.startswith("\xa0"):
            break
        body.append(line)
    out = []
    for r in csv.DictReader(io.StringIO("\n".join(body))):
        if (r.get("Asset Class") or "").strip() != "Equity":
            continue
        out.append({"ticker_raw": r.get("Ticker"), "cusip": _cusip(r.get("CUSIP")), "name": r.get("Name"),
                    "weight": _weight(r.get("Weight (%)"))})
    return out


def parse_vanguard(data: bytes) -> list[dict]:
    """Vanguard's portfolio-holding API: ``{"fund": {"entity": [{"ticker", "shortName", "cusip",
    "percentWeight", ...}]}}``. Lines without a ticker (cash, futures) are dropped."""
    body = json.loads(data)
    try:
        entities = body["fund"]["entity"]
    except (KeyError, TypeError) as e:
        raise base.SourceError("vanguard: no fund.entity list") from e
    return [{"ticker_raw": e.get("ticker"), "cusip": _cusip(e.get("cusip")), "name": e.get("shortName"),
             "weight": _weight(e.get("percentWeight"))} for e in entities or [] if (e or {}).get("ticker")]


def parse_nasdaq(data: bytes) -> list[dict]:
    """Nasdaq's list API: ``{"data": {"data": {"rows": [{"symbol", "companyName", ...}]}}}``."""
    body = json.loads(data)
    try:
        rows = body["data"]["data"]["rows"]
    except (KeyError, TypeError) as e:
        raise base.SourceError("nasdaq: no data.data.rows list") from e
    return [{"ticker_raw": r.get("symbol"), "cusip": None, "name": r.get("companyName"), "weight": None}
            for r in rows or [] if (r or {}).get("symbol")]


# ── The files ─────────────────────────────────────────────────────────────────────────────────────────────────

def _base(env: str, default: str) -> str:
    return (os.getenv(env) or default).strip().rstrip("/")


@dataclass(frozen=True)
class FundFile:
    index_slug: str
    provider: str
    fund: str
    ext: str
    parse: Callable[[bytes], list[dict]]

    def url(self) -> str:
        if self.provider == "ssga":
            return (_base("RC_SSGA_BASE", "https://www.ssga.com") + "/us/en/intermediary/library-content/products/"
                    f"fund-data/etfs/us/holdings-daily-us-en-{self.fund.lower()}.xlsx")
        if self.provider == "ishares":
            pid, slug = ISHARES_PRODUCTS[self.fund]
            return (_base("RC_ISHARES_BASE", "https://www.ishares.com")
                    + f"/us/products/{pid}/{slug}/1467271812596.ajax"
                    f"?fileType=csv&fileName={self.fund}_holdings&dataType=fund")
        if self.provider == "vanguard":
            return (_base("RC_VANGUARD_BASE", "https://investor.vanguard.com")
                    + f"/investment-products/etfs/profile/api/{self.fund}/portfolio-holding/stock")
        return _base("RC_NASDAQ_API_BASE", "https://api.nasdaq.com") + "/api/quote/list-type/nasdaq100"

    def raw_name(self) -> str:
        return f"{self.provider}-{self.fund}.{self.ext}"


FILES = (
    FundFile("sp500", "ssga", "SPY", "xlsx", parse_ssga),
    FundFile("sp500", "ishares", "IVV", "csv", parse_ishares),
    FundFile("sp500", "vanguard", "VOO", "json", parse_vanguard),
    FundFile("sp400", "ssga", "MDY", "xlsx", parse_ssga),
    FundFile("sp400", "ishares", "IJH", "csv", parse_ishares),
    FundFile("sp400", "vanguard", "IVOO", "json", parse_vanguard),
    FundFile("sp600", "ssga", "SPSM", "xlsx", parse_ssga),
    FundFile("sp600", "ishares", "IJR", "csv", parse_ishares),
    FundFile("sp600", "vanguard", "VIOO", "json", parse_vanguard),
    FundFile("ndx", "nasdaq", "NDX", "json", parse_nasdaq),
)
PROVIDERS = tuple(dict.fromkeys(f.provider for f in FILES))
_BY_RAW = {f.raw_name(): f for f in FILES}


def raw_dir() -> Path:
    return base.raw_root() / "funds" / "index_holdings"


def parquet_dir() -> Path:
    return base.parquet_root() / "funds" / "index_holdings"


def _meta(raw: Path) -> Path:
    return raw.with_name(raw.name + META_SUFFIX)


class FundHoldings:
    name = NAME
    kind = KIND
    providers = PROVIDERS
    tables: tuple[str, ...] = ()
    cadence = "daily"
    backfill_needs_from = False
    starts_by_hand = True             # not due until a first CLI run or Fetch built it (review should-fix 2)

    def default_as_of(self) -> date:
        return calendar.today_ny()

    def fetch(self, ctx: base.RunContext, as_of: date) -> list[Path]:
        out, failed = [], []
        for f in FILES:
            raw = raw_dir() / as_of.isoformat() / f.raw_name()
            if raw.exists():
                out.append(raw)
                continue
            try:
                r = base.http_get(ctx, f.provider, f.url(), headers={"Accept": "*/*"})
                f.parse(r.content)                                # refuse a file of the wrong shape before keeping it
            except Exception as e:  # noqa: BLE001 - one file down must not stop the others
                failed.append(f.raw_name())
                ctx.issue("source_failed", f.raw_name(), f"{type(e).__name__}: {e}"[:300])
                continue
            base.write_raw_json(_meta(raw), {"as_of": as_of.isoformat(), "fetched_at": time.time(),
                                             "run_id": ctx.run_id, "url": f.url().split("?")[0]})
            out.append(base.write_raw(raw, r.content))
        if failed:
            ctx.notes.append("failed files: " + ", ".join(failed))
            if len(failed) == len(FILES):
                raise base.SourceError("every fund file failed")
        return out

    def backfill(self, ctx: base.RunContext, start: date | None, end: date) -> None:
        # The files hold today's holdings only: a backfill is today's fetch. Membership history is WRDS's.
        self.to_parquet_all(ctx, self.fetch(ctx, end))

    def to_parquet(self, ctx: base.RunContext, raw: Path) -> list[Path]:
        return self.to_parquet_all(ctx, [raw])

    def to_parquet_all(self, ctx: base.RunContext, raws: list[Path]) -> list[Path]:
        by_day: dict[date, list[dict]] = {}
        for raw in raws:
            f = _BY_RAW.get(raw.name)
            if f is None:
                continue
            meta = base.read_raw_json(_meta(raw))
            as_of = date.fromisoformat(meta["as_of"])
            lines = f.parse(base.read_raw(raw))
            day = by_day.setdefault(as_of, [])
            for line in lines:
                day.append({"as_of": as_of, "index_slug": f.index_slug, "provider": f.provider, "fund": f.fund,
                            "ticker": norm_ticker(line["ticker_raw"]), "ticker_raw": line["ticker_raw"],
                            "cusip": line["cusip"], "name": line["name"], "weight": line["weight"],
                            "fetched_at": meta["fetched_at"], "run_id": meta["run_id"]})
        out = []
        for as_of, rows in sorted(by_day.items()):
            final = parquet_dir() / f"date={as_of.isoformat()}" / "part-0.parquet"
            if final.exists():                       # the day's other files, typed by an earlier run, stay in
                con = duckdb.connect()
                try:
                    keep = [dict(zip([c for c, _ in SCHEMA], r)) for r in con.execute(
                        f"SELECT * FROM read_parquet({base.sql_path(final)})").fetchall()]
                finally:
                    con.close()
                new = {(r["provider"], r["fund"]) for r in rows}
                rows = [r for r in keep if (r["provider"], r["fund"]) not in new] + rows
            con = duckdb.connect()
            try:
                base.rows_table(con, "rows", SCHEMA, rows)
                out.append(base.write_parquet(con, "SELECT * FROM rows ORDER BY index_slug, provider, ticker", final))
            finally:
                con.close()
        return out

    def build(self, ctx: base.RunContext, con: duckdb.DuckDBPyConnection) -> dict[str, int]:
        stage.create(con, "stg_fund_lines", SCHEMA, [f for _, f in stage.partitions(parquet_dir(), "date")])
        floors = " ".join(f"WHEN '{k}' THEN {v}" for k, v in FLOOR.items())
        cols = ", ".join(f'"{c}" {t}' for c, t in FILE_SCHEMA)
        con.execute(f"CREATE OR REPLACE TEMP TABLE stg_fund_files ({cols})")
        con.execute("INSERT INTO stg_fund_files SELECT as_of, index_slug, provider, fund, count(*)::INTEGER, "
                    f"count(*) >= CASE index_slug {floors} ELSE 1 END FROM stg_fund_lines "
                    "WHERE ticker IS NOT NULL OR cusip IS NOT NULL GROUP BY as_of, index_slug, provider, fund")
        return {}

    def check(self, ctx: base.RunContext, con: duckdb.DuckDBPyConnection) -> dict:
        latest = con.execute("SELECT max(as_of) FROM stg_fund_files").fetchone()[0]
        counted = stage.count(con, "stg_fund_files", f"as_of = DATE '{latest}' AND counted") if latest else 0
        # The members themselves are voted in spine/security's build (after matching to the SEC file), whose
        # index_membership is in this build already: the current members against the live file's.
        current = f"source = '{NAME}' AND valid_to IS NULL"
        added, removed = stage.compare_live(con, f"SELECT index_slug, ticker FROM index_membership WHERE {current}",
                                            "index_membership", ["index_slug", "ticker"], where=current)
        return {"dataset": NAME, "result": "changed" if added or removed else "unchanged",
                "source_as_of": (latest or ctx.as_of or self.default_as_of()).isoformat(),
                "rows": stage.count(con, "stg_fund_lines", f"as_of = DATE '{latest}'") if latest else 0,
                "added": added, "removed": removed, "coverage": round(counted / len(FILES), 4)}


DATASET = FundHoldings()
